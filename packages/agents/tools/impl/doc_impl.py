"""Document parsing tools: read/extract tables from xlsx, csv, pdf, docx, txt, json, md."""
import json
import io

from ...workspace import Workspace


def _format_table_summary(data: list[dict], max_rows: int = 5) -> str:
    if not data:
        return "  (empty)"
    cols = list(data[0].keys())
    lines = [f"  columns ({len(cols)}): {', '.join(cols)}", f"  total rows: {len(data)}"]
    for i, row in enumerate(data[:max_rows]):
        trimmed = {k: v for k, v in row.items() if k in cols[:8]}
        lines.append(f"  row {i}: {trimmed}")
    if len(data) > max_rows:
        lines.append(f"  ... and {len(data) - max_rows} more rows")
    return "\n".join(lines)


def _read_xlsx(raw: bytes, filename: str) -> str:
    try:
        import openpyxl
    except ImportError:
        return "需要 openpyxl: pip install openpyxl"

    wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    parts = [f"File: {filename}", f"Sheets ({len(wb.sheetnames)}): {', '.join(wb.sheetnames)}"]
    for sn in wb.sheetnames:
        ws = wb[sn]
        all_rows = list(ws.iter_rows(values_only=True))
        if not all_rows:
            parts.append(f"\n--- {sn} ---\n  (empty)")
            continue
        header: list[str] = []
        seen: set[str] = set()
        for i, h in enumerate(all_rows[0]):
            base = str(h) if h is not None else f"col_{i}"
            name = base
            c = 1
            while name in seen:
                name = f"{base}_{c}"
                c += 1
            seen.add(name)
            header.append(name)
        rows = [dict(zip(header, r)) for r in all_rows[1:] if any(v is not None for v in r)]
        parts.append(f"\n--- {sn} ---\n{_format_table_summary(rows)}")
    wb.close()
    return "\n".join(parts)


def _extract_xlsx(raw: bytes, sheet: str) -> list[dict]:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    sheet_names = wb.sheetnames

    if not sheet:
        target_sheets = [sheet_names[0]]
    else:
        target_sheets = [s for s in sheet_names if sheet.lower() in s.lower()]
        if not target_sheets:
            wb.close()
            raise ValueError(f"找不到 sheet '{sheet}'，可用: {sheet_names}")

    all_rows = []
    for sn in target_sheets:
        ws = wb[sn]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        header: list[str] = []
        seen: set[str] = set()
        for i, h in enumerate(rows[0]):
            base = str(h) if h is not None else f"col_{i}"
            name = base
            c = 1
            while name in seen:
                name = f"{base}_{c}"
                c += 1
            seen.add(name)
            header.append(name)
        data_rows = [dict(zip(header, r)) for r in rows[1:] if any(v is not None for v in r)]
        for r in data_rows:
            r["_sheet"] = sn
        all_rows.extend(data_rows)

    wb.close()
    return all_rows


def _read_csv(raw: bytes, _filename: str) -> str:
    import csv

    text = raw.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    rows = [dict(row) for row in reader]
    summary = _format_table_summary(rows)
    return f"File: {_filename}\n{summary}"


def _extract_csv(raw: bytes) -> list[dict]:
    import csv

    text = raw.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    return [dict(row) for row in reader]


def _read_pdf(raw: bytes, _filename: str) -> str:
    try:
        import pdfplumber
    except ImportError:
        return "需要 pdfplumber: pip install pdfplumber"

    parts = [f"File: {_filename}"]
    with pdfplumber.open(io.BytesIO(raw)) as pdf:
        parts.append(f"Pages: {len(pdf.pages)}")
        for i, page in enumerate(pdf.pages):
            try:
                text = page.extract_text() or ""
                tables = page.extract_tables()
                page_info = [f"\n--- Page {i+1} ---"]
                if text.strip():
                    page_info.append(f"Text ({len(text)} chars): {text[:500]}")
                if tables:
                    page_info.append(f"Tables found: {len(tables)}")
                    for ti, tbl in enumerate(tables[:3]):
                        page_info.append(f"  Table {ti+1}: {len(tbl)} rows x {len(tbl[0]) if tbl else 0} cols")
                parts.append("\n".join(page_info))
            except Exception:
                parts.append(f"\n--- Page {i+1} ---\n  (unable to parse)")
    return "\n".join(parts)


def _extract_pdf_tables(raw: bytes) -> list[dict]:
    try:
        import pdfplumber
    except ImportError:
        raise ImportError("需要 pdfplumber: pip install pdfplumber")

    all_rows = []
    with pdfplumber.open(io.BytesIO(raw)) as pdf:
        for pi, page in enumerate(pdf.pages):
            try:
                tables = page.extract_tables()
                for ti, tbl in enumerate(tables):
                    if not tbl or not tbl[0]:
                        continue
                    header = [str(c) if c else f"col_{ci}" for ci, c in enumerate(tbl[0])]
                    for row in tbl[1:]:
                        if any(c is not None for c in row):
                            row_dict = dict(zip(header, [str(c) if c else "" for c in row]))
                            row_dict["_page"] = pi + 1
                            row_dict["_table"] = ti + 1
                            all_rows.append(row_dict)
            except Exception:
                continue
    return all_rows


def _read_docx(raw: bytes, _filename: str) -> str:
    try:
        from docx import Document
    except ImportError:
        return "需要 python-docx: pip install python-docx"

    doc = Document(io.BytesIO(raw))
    parts = [f"File: {_filename}"]
    paras = [p.text for p in doc.paragraphs if p.text.strip()]
    parts.append(f"Paragraphs: {len(paras)}")
    if paras:
        text = "\n".join(paras[:50])
        parts.append(f"--- Text ---\n{text[:2000]}")
    tables = doc.tables
    parts.append(f"Tables: {len(tables)}")
    for ti, tbl in enumerate(tables[:5]):
        rows = [[cell.text for cell in row.cells] for row in tbl.rows]
        parts.append(f"\n--- Table {ti+1}: {len(rows)} rows x {len(rows[0]) if rows else 0} cols ---")
        for ri, row in enumerate(rows[:5]):
            parts.append(f"  row {ri}: {row[:6]}")
        if len(rows) > 5:
            parts.append(f"  ... and {len(rows) - 5} more rows")
    return "\n".join(parts)


def _extract_docx_tables(raw: bytes) -> list[dict]:
    from docx import Document

    doc = Document(io.BytesIO(raw))
    all_rows = []
    for ti, tbl in enumerate(doc.tables):
        rows = [[cell.text for cell in row.cells] for row in tbl.rows]
        if not rows:
            continue
        header = rows[0]
        for row in rows[1:]:
            row_dict = dict(zip(header, row))
            row_dict["_table"] = ti + 1
            all_rows.append(row_dict)
    return all_rows


# ── Public impl functions ──


def read_document_impl(ws: Workspace, path: str) -> str:
    """Read document content. Auto-detects format from file extension.
    Returns full text or a structured summary with sheet names, columns, row counts, and samples.
    """
    p = ws.resolve(path)
    if not p.exists():
        return json.dumps({"error": f"文件不存在: {path}"})

    lower = path.lower()
    raw = p.read_bytes()

    try:
        if lower.endswith((".xlsx", ".xls")):
            return _read_xlsx(raw, path)
        elif lower.endswith(".csv"):
            return _read_csv(raw, path)
        elif lower.endswith(".pdf"):
            return _read_pdf(raw, path)
        elif lower.endswith(".docx"):
            return _read_docx(raw, path)
        elif lower.endswith((".txt", ".md", ".json")):
            return p.read_text(encoding="utf-8", errors="replace")
        else:
            return p.read_text(encoding="utf-8", errors="replace")
    except Exception as e:
        return json.dumps({"error": f"读取文档失败: {e}"}, ensure_ascii=False)


def extract_document_tables_impl(ws: Workspace, path: str, sheet: str = "") -> str:
    """Extract structured table data from a document.
    For xlsx: each row is a dict with column names as keys.
    For csv: DictReader output.
    For pdf: extracted tables from all pages.
    For docx: extracted tables from the document.
    Returns JSON array of row dicts, or an error object.
    """
    p = ws.resolve(path)
    if not p.exists():
        return json.dumps({"error": f"文件不存在: {path}"})

    lower = path.lower()
    raw = p.read_bytes()

    try:
        if lower.endswith((".xlsx", ".xls")):
            rows = _extract_xlsx(raw, sheet)
        elif lower.endswith(".csv"):
            rows = _extract_csv(raw)
        elif lower.endswith(".pdf"):
            rows = _extract_pdf_tables(raw)
        elif lower.endswith(".docx"):
            rows = _extract_docx_tables(raw)
        else:
            return json.dumps({"error": f"不支持从此格式提取表格: {path}"})
        return json.dumps(rows, ensure_ascii=False, default=str)
    except Exception as e:
        return json.dumps({"error": f"提取表格失败: {e}"}, ensure_ascii=False)
